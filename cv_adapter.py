# ============================================================
# CV ADAPTER - Adapta un CV a un aviso de trabajo usando IA
# ============================================================

import re
import anthropic
import requests
from bs4 import BeautifulSoup
from io import BytesIO
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pdfplumber
from fpdf import FPDF


def extraer_texto_pdf(file_bytes: bytes) -> str:
    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        texto = ""
        for page in pdf.pages:
            texto += page.extract_text() or ""
    return texto.strip()


def extraer_texto_word(file_bytes: bytes) -> str:
    doc = Document(BytesIO(file_bytes))
    return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])


def _extraer_id_linkedin(url: str) -> str:
    match = re.search(r'/jobs/view/(\d+)', url)
    if match:
        return match.group(1)
    match = re.search(r'currentJobId=(\d+)', url)
    if match:
        return match.group(1)
    return ""


def scrapear_aviso(url: str) -> str:
    try:
        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0.0.0 Safari/537.36"
            )
        }

        # ── LinkedIn: usar endpoint público de guest ──────────
        if "linkedin.com" in url:
            job_id = _extraer_id_linkedin(url)
            if job_id:
                api_url = f"https://www.linkedin.com/jobs-guest/jobs/api/jobPosting/{job_id}"
                response = requests.get(api_url, headers=headers, timeout=15)
                if response.status_code == 200:
                    soup = BeautifulSoup(response.text, "html.parser")
                    for tag in soup(["script", "style"]):
                        tag.decompose()
                    lineas = [l.strip() for l in soup.get_text(separator="\n").splitlines() if l.strip()]
                    return "\n".join(lineas)[:8000]

        # ── Otros sitios: scraping general ───────────────────
        response = requests.get(url, headers=headers, timeout=15)
        soup = BeautifulSoup(response.text, "html.parser")
        for tag in soup(["script", "style", "nav", "header", "footer"]):
            tag.decompose()
        lineas = [l.strip() for l in soup.get_text(separator="\n").splitlines() if l.strip()]
        return "\n".join(lineas)[:8000]

    except Exception as e:
        print(f"❌ Error scrapeando aviso: {e}")
        return ""


def adaptar_cv(cv_texto: str, aviso_texto: str) -> str:
    client = anthropic.Anthropic()

    prompt = f"""Eres un experto en recursos humanos y en optimización de CVs para sistemas ATS \
(Applicant Tracking Systems).

Tu tarea es adaptar el CV que te voy a pasar para que se ajuste mejor al aviso de trabajo, \
siguiendo estas reglas ESTRICTAS:

REGLAS:
1. NO inventar ni agregar información falsa. Solo trabajar con lo que está en el CV original.
2. NO alucinar experiencias, habilidades ni logros que no estén en el CV.
3. SÍ reorganizar y reformular el contenido existente para destacar lo más relevante.
4. SÍ incorporar palabras clave del aviso donde sean verdaderas y aplicables.
5. SÍ destacar habilidades blandas (comunicación, liderazgo, aprendizaje, adaptabilidad) \
cuando no haya coincidencia técnica exacta.
6. SÍ resaltar habilidades parciales: si el aviso pide A+B y el candidato tiene A, \
destacar A y mencionar disposición a aprender B.
7. Mantener la estructura general del CV original.
8. Devolver el CV adaptado completo, listo para usar.

AVISO DE TRABAJO:
{aviso_texto}

CV ORIGINAL:
{cv_texto}

Devolvé el CV adaptado completo. No incluyas explicaciones ni comentarios fuera del CV."""

    message = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=4096,
        messages=[{"role": "user", "content": prompt}]
    )

    return message.content[0].text


def _agregar_parrafo_con_negrita(doc, texto: str, justificado: bool = True):
    p = doc.add_paragraph()
    if justificado:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    partes = re.split(r'(\*\*[^*]+\*\*)', texto)
    for parte in partes:
        if parte.startswith('**') and parte.endswith('**'):
            run = p.add_run(parte[2:-2])
            run.bold = True
        elif parte:
            p.add_run(parte)
    return p


def _es_fila_tabla(linea: str) -> bool:
    return linea.startswith('|') and linea.endswith('|')

def _es_separador_tabla(linea: str) -> bool:
    return _es_fila_tabla(linea) and all(c in '-| :' for c in linea)

def _celdas(linea: str) -> list:
    return [c.strip() for c in linea.strip('|').split('|')]

def _agregar_tabla_word(doc, filas: list):
    from docx.shared import RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    if not filas:
        return
    cols = len(_celdas(filas[0]))
    tabla = doc.add_table(rows=0, cols=cols)
    tabla.style = 'Table Grid'

    for i, fila in enumerate(filas):
        celdas = _celdas(fila)
        row = tabla.add_row()
        for j, celda in enumerate(celdas):
            cell = row.cells[j]
            cell.text = celda
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph('')


def generar_word(texto: str) -> bytes:
    doc = Document()

    estilo = doc.styles['Normal']
    estilo.font.name = 'Calibri'
    estilo.font.size = Pt(11)

    lineas = texto.split("\n")
    lineas_vacias_consecutivas = 0
    i = 0

    while i < len(lineas):
        linea_strip = lineas[i].strip()

        if linea_strip in ('---', '***', '___'):
            i += 1
            continue

        if not linea_strip:
            lineas_vacias_consecutivas += 1
            if lineas_vacias_consecutivas <= 1:
                doc.add_paragraph('')
            i += 1
            continue

        lineas_vacias_consecutivas = 0

        # ── Detectar tabla Markdown ───────────────────────────
        if _es_fila_tabla(linea_strip):
            filas_tabla = []
            while i < len(lineas) and _es_fila_tabla(lineas[i].strip()):
                if not _es_separador_tabla(lineas[i].strip()):
                    filas_tabla.append(lineas[i].strip())
                i += 1
            _agregar_tabla_word(doc, filas_tabla)
            continue

        if linea_strip.startswith('# '):
            doc.add_heading(linea_strip[2:], level=1)
        elif linea_strip.startswith('## '):
            doc.add_heading(linea_strip[3:], level=2)
        elif linea_strip.startswith('### '):
            doc.add_heading(linea_strip[4:], level=3)
        else:
            _agregar_parrafo_con_negrita(doc, linea_strip)

        i += 1

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _limpiar_para_pdf(texto: str) -> str:
    reemplazos = {
        '\u2013': '-', '\u2014': '-',
        '\u2018': "'", '\u2019': "'",
        '\u201c': '"', '\u201d': '"',
        '\u2022': '-', '\u2026': '...',
        '\u00b7': '-', '\u2015': '-',
        '\u00e2': 'a', '\u20ac': 'EUR',
    }
    for char, reemplazo in reemplazos.items():
        texto = texto.replace(char, reemplazo)
    return texto.encode('latin-1', errors='ignore').decode('latin-1')


def _pdf_linea_con_negrita(pdf, texto: str, h: int = 7):
    partes = re.split(r'(\*\*[^*]+\*\*)', texto)
    for parte in partes:
        if parte.startswith('**') and parte.endswith('**'):
            pdf.set_font("Helvetica", style="B", size=11)
            pdf.write(h, parte[2:-2])
        elif parte:
            pdf.set_font("Helvetica", size=11)
            pdf.write(h, parte)
    pdf.ln(h)


def generar_pdf(texto: str) -> bytes:
    pdf = FPDF()
    pdf.set_margins(20, 20, 20)
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=20)

    lineas_vacias_consecutivas = 0

    for linea in texto.split("\n"):
        linea_strip = _limpiar_para_pdf(linea.strip())

        if linea_strip in ('---', '***', '___'):
            continue

        if not linea_strip:
            lineas_vacias_consecutivas += 1
            if lineas_vacias_consecutivas <= 1:
                pdf.ln(4)
            continue

        lineas_vacias_consecutivas = 0

        try:
            if linea_strip.startswith('# '):
                pdf.set_font("Helvetica", style="B", size=16)
                pdf.multi_cell(0, 9, linea_strip[2:])
            elif linea_strip.startswith('## '):
                pdf.set_font("Helvetica", style="B", size=13)
                pdf.multi_cell(0, 8, linea_strip[3:])
            elif linea_strip.startswith('### '):
                pdf.set_font("Helvetica", style="B", size=11)
                pdf.multi_cell(0, 7, linea_strip[4:])
            else:
                _pdf_linea_con_negrita(pdf, linea_strip)
        except Exception:
            pdf.ln(4)

    return bytes(pdf.output())
